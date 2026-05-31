"""
insert_figures.py — Inserts Figures 1-8 into the MDPI manuscript docx at their caption locations.

Use after build_mdpi.py or when starting from a docx that has figure captions but no embedded
figure images. This addresses Round 7 Blocker #1 (Figures 1-8 not embedded in document body).

Strategy: locate paragraphs starting with "Fig. N." (case-sensitive), and insert a centered
image paragraph immediately BEFORE the caption paragraph.

Usage (workstation):
    cd radiation-mapping/ukedo_river2
    python tools/insert_figures.py

Inputs (auto-detected, paths can be edited at top of script):
    docs/manuscript_mdpi.docx          # source docx with captions but no figures
    data/results_v81/Figure_{1..8}.png # figure files

Output:
    docs/manuscript_mdpi.docx          # in-place update (back up first if you want)
    docs/manuscript_mdpi_with_figures.docx (alternative output if NEW_OUTPUT below set True)
"""
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    print("ERROR: python-docx not installed. Run: pip install python-docx")
    sys.exit(1)

# =========================
# Path resolution (portable)
# =========================
HERE = Path(__file__).resolve().parent
PROJECT_ROOT = HERE.parent  # tools/.. = ukedo_river2/

DOCX_IN  = PROJECT_ROOT / 'docs' / 'manuscript_mdpi.docx'
FIG_DIR  = PROJECT_ROOT / 'data' / 'results_v81'
NEW_OUTPUT = False  # True → save as manuscript_mdpi_with_figures.docx; False → overwrite

# =========================
# Configuration
# =========================
FIGURE_NUMBERS = list(range(1, 9))  # Figures 1-8
FIGURE_WIDTH_INCHES = 6.5            # MDPI Remote Sensing column-width compatible
CAPTION_PREFIX_VARIANTS = [
    "Fig. {n}.",       # Most common
    "Figure {n}.",     # Alternative spelling
    "**Fig. {n}.**",   # If bold markdown preserved
]


def find_caption_paragraph(doc, fig_n):
    """Find the paragraph whose text begins with 'Fig. N.' or variant for figure fig_n."""
    expected_prefixes = [v.format(n=fig_n) for v in CAPTION_PREFIX_VARIANTS]
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        for prefix in expected_prefixes:
            if text.startswith(prefix):
                return i, para
    return None, None


def insert_image_before(doc, anchor_para, image_path, width_inches):
    """Insert a centered image paragraph immediately before anchor_para."""
    new_para = anchor_para.insert_paragraph_before()
    new_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = new_para.add_run()
    run.add_picture(str(image_path), width=Inches(width_inches))


def main():
    # Validate inputs
    if not DOCX_IN.exists():
        sys.exit(f"ERROR: docx not found at {DOCX_IN}")
    if not FIG_DIR.exists():
        sys.exit(f"ERROR: figure directory not found at {FIG_DIR}")

    missing_figs = [n for n in FIGURE_NUMBERS if not (FIG_DIR / f'Figure_{n}.png').exists()]
    if missing_figs:
        print(f"WARNING: Missing figure files: {missing_figs}")
        print(f"Continuing — present figures will still be inserted.")

    print(f"Loading {DOCX_IN}...")
    doc = Document(str(DOCX_IN))

    # Diagnostic: count existing drawings
    from docx.oxml.ns import qn
    body_xml = doc.element.body
    existing_drawings = body_xml.findall('.//' + qn('w:drawing'))
    print(f"  Existing <w:drawing> elements in body: {len(existing_drawings)}")

    inserted = 0
    skipped = 0
    for fig_n in FIGURE_NUMBERS:
        fig_path = FIG_DIR / f'Figure_{fig_n}.png'
        if not fig_path.exists():
            print(f"  Fig {fig_n}: SKIP (file not found)")
            skipped += 1
            continue

        para_idx, anchor = find_caption_paragraph(doc, fig_n)
        if anchor is None:
            print(f"  Fig {fig_n}: SKIP (caption paragraph not found in docx)")
            skipped += 1
            continue

        insert_image_before(doc, anchor, fig_path, FIGURE_WIDTH_INCHES)
        print(f"  Fig {fig_n}: INSERTED before paragraph [{para_idx}] (caption: '{anchor.text[:60]}...')")
        inserted += 1

    # Save
    if NEW_OUTPUT:
        out_path = PROJECT_ROOT / 'docs' / 'manuscript_mdpi_with_figures.docx'
    else:
        out_path = DOCX_IN
    doc.save(str(out_path))
    print(f"\nSaved {out_path}")
    print(f"Inserted: {inserted} figure(s)")
    print(f"Skipped:  {skipped} figure(s)")

    if inserted > 0:
        # Verify post-insertion
        doc2 = Document(str(out_path))
        body_xml2 = doc2.element.body
        new_drawings = body_xml2.findall('.//' + qn('w:drawing'))
        print(f"\nVerification: <w:drawing> elements after save: {len(new_drawings)}")
        print(f"  (expected: {len(existing_drawings) + inserted})")


if __name__ == '__main__':
    main()
